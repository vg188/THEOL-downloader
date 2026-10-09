"""Optional fixture regeneration; tests read checked-in files, not these dependencies.
Creates synthetic documents only. No account, course data, or network is used.
Requires PyMuPDF, python-docx, python-pptx, openpyxl, Pillow.
Legacy conversion is a separate, optional local Office step.
"""
from pathlib import Path
import fitz
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from PIL import Image

out = Path(__file__).resolve().parent
pdf = fitz.open()
page = pdf.new_page()
page.insert_text((72, 100), "BUCT download integrity fixture - PDF - 20261008", fontsize=14)
pdf.save(str(out / "sample.pdf"))
pdf.close()
other = fitz.open()
other.new_page().insert_text((72, 100), "BUCT alternate PDF - distinct content, same saved name", fontsize=14)
other.save(str(out / "sample-alt.pdf"))
other.close()
doc = Document()
doc.core_properties.author = "BUCT synthetic fixture"
doc.core_properties.last_modified_by = "BUCT synthetic fixture"
doc.add_heading("Original file integrity fixture", 0)
doc.add_paragraph("BUCT-fixture-docx-20261008")
doc.add_paragraph("中文内容、两个  空格、数学符号：α + β = γ")
doc.save(out / "sample.docx")
prs = Presentation()
prs.core_properties.author = "BUCT synthetic fixture"
prs.core_properties.last_modified_by = "BUCT synthetic fixture"
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "BUCT-fixture-pptx-20261008"
slide.placeholders[1].text = "原始幻灯片内容\n文件名、格式、字节保持对应"
prs.save(out / "sample.pptx")
workbook = Workbook()
workbook.properties.creator = "BUCT synthetic fixture"
workbook.properties.lastModifiedBy = "BUCT synthetic fixture"
sheet = workbook.active
sheet.title = "Integrity"
sheet.append(["BUCT-fixture-xlsx-20261008", "原始单元格内容"])
sheet.append(["value", 42])
sheet.append(["sum", "=SUM(1,2)"])
workbook.save(out / "sample.xlsx")
image = Image.new("RGB", (24, 16), (32, 97, 204))
image.save(out / "sample.png")
image.save(out / "sample.jpg", quality=90)
print("Created PDF/DOCX/PPTX/XLSX/PNG/JPEG synthetic fixtures")
